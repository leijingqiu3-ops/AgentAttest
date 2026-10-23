from __future__ import annotations

import argparse
import csv
import json
import os
import sys
import zipfile
from pathlib import Path

import openpyxl

from .config import ROOT, get_client, get_settings
from .models import Verdict
from .report import render_markdown
from .recompute import recompute
from .runner import load_cases, run_suite
from .target.agent import FAULT_MODES

DEFAULT_XLSX = Path(os.environ.get("AUDIT_BADCASE_XLSX", str(ROOT / "data" / "badcases.xlsx")))


def cmd_build_cases(args: argparse.Namespace) -> None:
    from scripts.build_seed_cases import main as build_main

    build_main()


def _suggest(description: str, record_type: str) -> tuple[str, str, str]:
    text = f"{record_type} {description}"
    rules = [
        (("字段", "手册", "样例表", "不一致"), "doc_field_mismatch", "可转换：补预期字段即可"),
        (("图表", "专业比例", "类别"), "chart_mismatch", "可转换：补预期类别数即可"),
        (("依赖", "ModuleNotFound", "库", "环境"), "env_dependency", "可转换：补精确包名即可"),
        (("权限", "越权", "其他课程"), "permission", "可转换：补角色与课程范围"),
        (("重复", "工单", "又出现"), "duplicate_ticket", "可转换：补原工单号"),
        (("退课", "退款", "期限"), "policy_exception", "可转换：补政策条款"),
        (("成绩", "缺失", "空值", "数据"), "missing_grade", "可转换：补学生与课程"),
        (("慢", "性能", "超时", "卡顿"), "performance", "需模拟环境：当前靶子未覆盖"),
        (("点击", "界面", "按钮", "交互", "刷新"), "interaction", "需模拟环境：当前靶子未覆盖"),
    ]
    for keywords, family, note in rules:
        if any(keyword in text for keyword in keywords):
            return family, note.split("：")[0], note.split("：", 1)[1] if "：" in note else ""
    return "other", "需人工判断", "描述不足，需补预期与实际"


def cmd_import_badcases(args: argparse.Namespace) -> None:
    source = Path(args.xlsx)
    if not source.exists():
        raise SystemExit(f"找不到 xlsx: {source}")
    workbook = openpyxl.load_workbook(source, data_only=True)
    sheet = workbook[workbook.sheetnames[1]]
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    records = [row for row in rows if row and row[0] not in (None, "")]
    images = 0
    with zipfile.ZipFile(source) as archive:
        images = len([name for name in archive.namelist() if name.startswith("xl/media/")])

    output = Path(args.out)
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["编号", "日期", "任务", "记录类型", "原描述", "建议场景", "建议失败类", "可转换性", "转写备注"])
        for row in records:
            number, day, task, record_type, description, _note = (list(row) + [None] * 6)[:6]
            family, feasibility, note = _suggest(str(description or ""), str(record_type or ""))
            writer.writerow([number, day, task, record_type, description, family, family, feasibility, note])
    print(f"imported {len(records)} records, {images} embedded screenshots -> {output}")


def cmd_run(args: argparse.Namespace) -> None:
    cases = load_cases(args.suite)
    if args.families:
        wanted = {item.strip() for item in args.families.split(",") if item.strip()}
        cases = [case for case in cases if case.family in wanted]
    if args.one_per_family:
        seen: set[str] = set()
        picked = []
        for case in cases:
            if case.family not in seen:
                seen.add(case.family)
                picked.append(case)
        cases = picked
    if args.limit:
        cases = cases[: args.limit]
    if not cases:
        raise SystemExit("没有匹配的用例")
    faults = [item.strip() for item in args.faults.split(",") if item.strip()]
    unknown = [item for item in faults if item not in FAULT_MODES]
    if unknown:
        raise SystemExit(f"未知故障模式: {unknown}; 可选 {FAULT_MODES}")
    settings = get_settings()
    client = get_client(settings)
    factory = None
    if args.agent == "langgraph":
        from .adapters import LangGraphReActAgent

        factory = lambda s: LangGraphReActAgent(s)  # noqa: E731
        if faults != ["none"]:
            raise SystemExit("LangGraph 适配器只支持 --faults none（真实 Agent 不做注入）")
    summary = run_suite(
        cases, faults, args.k, args.out, settings, client,
        agent_factory=factory, agent_name=args.agent,
    )
    verdicts = [
        Verdict.model_validate_json(line)
        for line in (Path(args.out) / "verdicts.jsonl").read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]
    report_path = Path(args.out) / "report.md"
    render_markdown(summary, verdicts, report_path)
    print(json.dumps(summary["none"], ensure_ascii=False, indent=2))
    print(f"report: {report_path}")


def cmd_report(args: argparse.Namespace) -> None:
    run_dir = Path(args.run)
    summary = json.loads((run_dir / "summary.json").read_text(encoding="utf-8"))
    verdicts = [
        Verdict.model_validate_json(line)
        for line in (run_dir / "verdicts.jsonl").read_text(encoding="utf-8").splitlines()
        if line.strip()
    ]
    path = render_markdown(summary, verdicts, run_dir / "report.md")
    print(path)


def cmd_recompute(args: argparse.Namespace) -> None:
    summary = recompute(args.run, args.suite)
    print(json.dumps({"none": summary["none"], "faults": summary["faults"]}, ensure_ascii=False, indent=2))


def cmd_selfcheck(args: argparse.Namespace) -> None:
    from .selfcheck import run_selfcheck

    result = run_selfcheck()
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if not result["passed"]:
        raise SystemExit(1)


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(prog="audit_agent", description="Agent 验收/审计 Agent")
    sub = parser.add_subparsers(dest="command", required=True)

    build = sub.add_parser("build-cases", help="生成 40 条种子用例")
    build.set_defaults(func=cmd_build_cases)

    imp = sub.add_parser("import-badcases", help="把平台 badcase xlsx 转成候选用例 CSV")
    imp.add_argument("--xlsx", default=str(DEFAULT_XLSX))
    imp.add_argument("--out", default=str(ROOT / "cases" / "from_badcases.csv"))
    imp.set_defaults(func=cmd_import_badcases)

    run = sub.add_parser("run", help="运行审计套件")
    run.add_argument("--suite", default=str(ROOT / "cases" / "seed_cases.jsonl"))
    run.add_argument("--faults", default="none,wrong_arg,skip_clarify,context_drop,policy_ignore,pii_leak,early_stop")
    run.add_argument("--k", type=int, default=1)
    run.add_argument("--limit", type=int, default=0)
    run.add_argument("--families", default="")
    run.add_argument("--one-per-family", action="store_true")
    run.add_argument("--agent", choices=["synth", "langgraph"], default="synth")
    run.add_argument("--out", default=str(ROOT / "reports" / "latest"))
    run.set_defaults(func=cmd_run)

    report = sub.add_parser("report", help="从已有运行结果重新生成报告")
    report.add_argument("--run", default=str(ROOT / "reports" / "latest"))
    report.set_defaults(func=cmd_report)

    recalc = sub.add_parser("recompute", help="离线重算断言与归因，不重复调用 API")
    recalc.add_argument("--run", default=str(ROOT / "reports" / "latest"))
    recalc.add_argument("--suite", default=str(ROOT / "cases" / "seed_cases.jsonl"))
    recalc.set_defaults(func=cmd_recompute)

    check = sub.add_parser("selfcheck", help="用 known-answer 与变异体验证 oracle 自身")
    check.set_defaults(func=cmd_selfcheck)

    args = parser.parse_args(argv)
    args.func(args)


if __name__ == "__main__":
    main()
